import pandas as pd
import numpy as np
import os
from docx import Document

def main():
    doc_path = r"G:\Gauranshi\IIITD-Research\Atlas_opt\Opt-exp-res\ATLAS OPTIMIZER EXPERIMENTS_060926.docx"
    doc = Document(doc_path)
    
    doc.add_page_break()
    doc.add_heading("FINAL UPDATED RESULTS (Including New Muon/Muon-SAM Fix)", level=1)
    
    runs = [
        ('Experiment 1: CIFAR-10 Image Classification', 'final_logs/run1_cifar10_logs.csv'),
        ('Experiment 2: nanoGPT Continual Pre-Training', 'final_logs/run2_nanogpt_logs.csv'),
        ('Experiment 3: Pythia-70M Continual Pre-Training', 'final_logs/run3_pythia_logs.csv')
    ]
    
    for title, csv_path in runs:
        if not os.path.exists(csv_path): continue
        df = pd.read_csv(csv_path)
        
        doc.add_heading(title, level=2)
        
        has_ppl = 'cifar' not in title.lower()
        cols = 6 if has_ppl else 5
        
        table = doc.add_table(rows=1, cols=cols)
        hdr = table.rows[0].cells
        hdr[0].text = 'Optimizer'
        hdr[1].text = 'Final Train Loss'
        hdr[2].text = 'Final Train Acc'
        hdr[3].text = 'Final Val Loss'
        hdr[4].text = 'Final Val Acc'
        if has_ppl:
            hdr[5].text = 'Final Val PPL'
            
        for cell in hdr:
            cell.paragraphs[0].runs[0].bold = True
            
        for opt in ['atlas_random', 'atlas_raw', 'atlas', 'adam', 'muon', 'muon_sam', 'sgd']:
            opt_lower = opt.lower()
            opt_df = df[df['optimizer'].str.lower() == opt_lower]
            if opt_df.empty: continue
            
            final_row = opt_df.iloc[-1]
            t_loss = f"{final_row['train_loss']:.4f}"
            t_acc = f"{final_row['train_accuracy']:.2f}%"
            v_loss = f"{final_row['val_loss']:.4f}"
            v_acc = f"{final_row['val_accuracy']:.2f}%"
            
            display_name = opt.upper()
            if display_name == 'MUON_SAM': display_name = 'Muon-SAM'
            elif display_name == 'ADAM': display_name = 'AdamW'
            elif display_name == 'SGD': display_name = 'SGD'
            elif display_name == 'MUON': display_name = 'Muon'
            elif display_name == 'ATLAS': display_name = 'Atlas'
            elif display_name == 'ATLAS_RAW': display_name = 'Atlas Raw'
            elif display_name == 'ATLAS_RANDOM': display_name = 'Atlas Random'
            
            row = table.add_row().cells
            row[0].text = display_name
            row[1].text = t_loss
            row[2].text = t_acc
            row[3].text = v_loss
            row[4].text = v_acc
            if has_ppl:
                row[5].text = f"{np.exp(final_row['val_loss']):.2f}"
                
        doc.add_paragraph()
        
    doc.save(doc_path)
    print("Word document updated successfully!")

if __name__ == '__main__':
    main()
