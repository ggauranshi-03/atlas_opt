import os
from docx import Document

def main():
    doc_path = r"G:\Gauranshi\IIITD-Research\Atlas_opt\Opt-exp-res\ATLAS OPTIMIZER EXPERIMENTS_060926.docx"
    doc = Document(doc_path)
    
    doc.add_page_break()
    doc.add_heading("Final Validation Perplexity Results", level=1)
    
    # NanoGPT
    doc.add_heading("Experiment 2: nanoGPT Continual Pre-Training", level=2)
    table2 = doc.add_table(rows=1, cols=2)
    hdr2 = table2.rows[0].cells
    hdr2[0].text, hdr2[1].text = 'Optimizer', 'Final Val Perplexity'
    hdr2[0].paragraphs[0].runs[0].bold, hdr2[1].paragraphs[0].runs[0].bold = True, True
    
    res2 = [
        ('Atlas Random', '42.16'), ('Atlas Raw', '43.61'), ('Atlas', '57.00'),
        ('AdamW', '85.17'), ('Muon', '193.48'), ('Muon-SAM', '332.40'), ('SGD (Fixed)', '3906.18')
    ]
    for opt, ppl in res2:
        row = table2.add_row().cells
        row[0].text, row[1].text = opt, ppl
    doc.add_paragraph()
    
    # Pythia
    doc.add_heading("Experiment 3: Pythia-70M Continual Pre-Training", level=2)
    table3 = doc.add_table(rows=1, cols=2)
    hdr3 = table3.rows[0].cells
    hdr3[0].text, hdr3[1].text = 'Optimizer', 'Final Val Perplexity'
    hdr3[0].paragraphs[0].runs[0].bold, hdr3[1].paragraphs[0].runs[0].bold = True, True
    
    res3 = [
        ('Atlas Random', '55.18'), ('Atlas Raw', '55.92'), ('Atlas', '64.15'),
        ('AdamW', '136.40'), ('Muon', '379.22'), ('Muon-SAM', '584.19'), ('SGD (Fixed)', '3161.76')
    ]
    for opt, ppl in res3:
        row = table3.add_row().cells
        row[0].text, row[1].text = opt, ppl
        
    doc.save(doc_path)

if __name__ == '__main__':
    main()
